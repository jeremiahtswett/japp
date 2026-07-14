from docx import Document

from japp.tailoring.docx_render import render_resume


def _all_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def test_render_resume_produces_expected_content(tmp_path, sample_corpus, sample_validated):
    out_path = tmp_path / "resume.docx"
    render_resume(sample_corpus, sample_validated, out_path)
    assert out_path.exists()

    doc = Document(str(out_path))
    text = _all_text(doc)

    assert "Sam Test" in text
    assert "sam@test.com" in text
    assert "Built a Python microservice handling 10k orders/day via REST APIs." in text
    assert "Experienced software engineer skilled in Python and SQL." in text
    assert "Acme Corp - Software Engineer" in text
    assert "State University" in text
    assert "Python, SQL, Flask" in text

    # excluded section must not appear
    assert "Coding Club" not in text
    # never use tables/text boxes (spec Stage 3 requirement)
    assert doc.tables == []


def test_render_resume_omits_summary_when_absent(tmp_path, sample_corpus, sample_validated):
    validated = dict(sample_validated, summary_line=None)
    out_path = tmp_path / "resume.docx"
    render_resume(sample_corpus, validated, out_path)
    text = _all_text(Document(str(out_path)))
    assert "Experienced software engineer" not in text


def test_render_resume_creates_parent_dir(tmp_path, sample_corpus, sample_validated):
    out_path = tmp_path / "nested" / "dir" / "resume.docx"
    render_resume(sample_corpus, sample_validated, out_path)
    assert out_path.exists()


def test_education_with_empty_detail_has_no_stray_dash(tmp_path, sample_corpus, sample_validated):
    corpus = dict(sample_corpus)
    corpus["education"] = corpus["education"] + [
        {"id": "edu-2", "institution": "MIT (Cross-Registered)", "detail": "",
         "dates": "2022-Present", "highlights": []},
    ]
    out_path = tmp_path / "resume.docx"
    render_resume(corpus, sample_validated, out_path)
    text = _all_text(Document(str(out_path)))
    assert "MIT (Cross-Registered) (2022-Present)" in text
    assert "MIT (Cross-Registered) - " not in text
