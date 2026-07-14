"""Render the tailored resume as a clean, single-column .docx.

ADR 0005: python-docx over HTML->PDF. No tables, no text boxes, no multi-column
layout - deliberately standardized rather than pixel-matching the master's
original design, per spec Stage 3's "clean, single-column, machine-parseable"
requirement.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _section_heading_text(sec: dict) -> str:
    return sec["name"] if not sec.get("title") else f"{sec['name']} - {sec['title']}"


def render_resume(corpus: dict, validated: dict, output_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin = section.bottom_margin = Inches(0.5)

    contact = corpus.get("contact", {})
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(contact.get("name") or "")
    name_run.bold = True
    name_run.font.size = Pt(18)

    contact_bits = [contact.get("location", ""), contact.get("email", ""), *contact.get("links", [])]
    contact_line = " | ".join(b for b in contact_bits if b)
    if contact_line:
        cp = doc.add_paragraph(contact_line)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if validated.get("summary_line"):
        sp = doc.add_paragraph()
        run = sp.add_run(validated["summary_line"])
        run.italic = True

    for sec in validated["sections"]:
        doc.add_heading(_section_heading_text(sec), level=2)
        meta = " | ".join(b for b in (sec.get("location", ""), sec.get("dates", "")) if b)
        if meta:
            mp = doc.add_paragraph(meta)
            for run in mp.runs:
                run.italic = True
                run.font.size = Pt(9)
        for bullet in sec["bullets"]:
            doc.add_paragraph(bullet["tailored_text"], style="List Bullet")

    if corpus.get("education"):
        doc.add_heading("Education", level=2)
        for edu in corpus["education"]:
            line = f"{edu['institution']} - {edu['detail']}"
            if edu.get("dates"):
                line += f" ({edu['dates']})"
            doc.add_paragraph(line)
            for highlight in edu.get("highlights", []):
                doc.add_paragraph(highlight, style="List Bullet")

    if corpus.get("skills"):
        doc.add_heading("Skills", level=2)
        doc.add_paragraph(", ".join(corpus["skills"]))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
